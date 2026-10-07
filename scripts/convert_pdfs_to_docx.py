import os
import sys
import json
import time
import re
import warnings
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor

warnings.filterwarnings("ignore")
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

BASE_DOCS_DIR = "Documents utiles"
PROGRESS_FILE = "conversion_progress.json"

def to_long_path(p):
    """Supporte les chemins longs (> 260 caractères) sous Windows."""
    ap = os.path.abspath(str(p))
    if os.name == 'nt' and not ap.startswith('\\\\?\\'):
        return '\\\\?\\' + ap
    return ap

def sanitize_xml(text):
    """Supprime les caractères de contrôle non compatibles XML."""
    if not text:
        return ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', text)

def convert_with_pymupdf_fast(pdf_path, docx_path):
    """Extraction haute vitesse et structurée en DOCX via PyMuPDF."""
    import docx
    import pymupdf
    
    doc = pymupdf.open(to_long_path(pdf_path))
    out_doc = docx.Document()
    
    for i, page in enumerate(doc):
        text = page.get_text()
        if text.strip():
            p_head = out_doc.add_paragraph()
            p_head.add_run(f"=== Page {i + 1} ===").bold = True
            
            for line in text.split("\n"):
                clean_line = sanitize_xml(line.strip())
                if clean_line:
                    out_doc.add_paragraph(clean_line)
                    
    out_doc.save(to_long_path(docx_path))
    doc.close()
    return True

def convert_with_pdf2docx_worker(pdf_str, docx_str):
    """Convertit avec pdf2docx dans un sous-processus isolé."""
    from pdf2docx import Converter
    cv = Converter(to_long_path(pdf_str))
    try:
        cv.convert(to_long_path(docx_str), start=0, end=None)
        return True, None
    except Exception as e:
        return False, str(e)
    finally:
        try:
            cv.close()
        except Exception:
            pass

def get_pdf_page_count(pdf_path):
    """Retourne rapidement le nombre de pages d'un PDF."""
    try:
        import pymupdf
        doc = pymupdf.open(to_long_path(pdf_path))
        cnt = len(doc)
        doc.close()
        return cnt
    except Exception:
        return 10

def update_progress(data):
    try:
        with open(PROGRESS_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def main():
    print("=" * 75)
    print("🚀 PROCURA - CONVERSION DES NOUVEAUX DOCUMENTS PDF -> WORD (.DOCX)")
    print("📌 Recherche automatique de tous les PDFs sans équivalent DOCX")
    print("📌 Conservation simultanée des PDF et DOCX dans chaque dossier")
    print("=" * 75)

    if not os.path.exists(BASE_DOCS_DIR):
        print(f"❌ {BASE_DOCS_DIR} introuvable.")
        return

    # Scanner récursivement tout le dossier Documents utiles
    pdf_list = []
    for root, dirs, files in os.walk(BASE_DOCS_DIR):
        files_lower = {f.lower(): f for f in files}
        for f in files:
            if f.lower().endswith(".pdf") and not f.startswith("~$") and f != ".pdf":
                stem = f.rsplit(".", 1)[0].lower()
                docx_name = stem + ".docx"
                if docx_name not in files_lower:
                    pdf_p = Path(root) / f
                    try:
                        # Vérifier que le fichier n'est pas vide (0 octet)
                        if os.path.getsize(to_long_path(pdf_p)) > 0:
                            docx_p = Path(root) / (Path(f).stem + ".docx")
                            pdf_list.append((pdf_p, docx_p))
                    except Exception:
                        pass

    total = len(pdf_list)
    print(f"\n📋 Nouveaux documents PDF à convertir : {total}\n")

    if total == 0:
        print("🎉 Aucun nouveau PDF à convertir : tous ont déjà leur équivalent .docx !")
        update_progress({"status": "completed", "total": 0, "processed": 0, "success": 0, "errors": 0})
        return

    success_count = 0
    start_time = time.time()

    progress_data = {
        "status": "running",
        "total": total,
        "processed": 0,
        "success": 0,
        "errors": 0,
        "current_file": "",
        "start_time": time.strftime("%Y-%m-%d %H:%M:%S")
    }
    update_progress(progress_data)

    for idx, (pdf_p, docx_p) in enumerate(pdf_list, 1):
        rel_folder = pdf_p.parent.name
        rel_name = pdf_p.name
        percent = (idx / total) * 100
        
        page_count = get_pdf_page_count(pdf_p)
        print(f"[{idx}/{total}] ({percent:.1f}%) [{rel_folder}] 📄 {rel_name} ({page_count}p) ...", end=" ", flush=True)

        progress_data["processed"] = idx
        progress_data["current_file"] = f"[{rel_folder}] {rel_name}"
        update_progress(progress_data)

        t0 = time.time()
        success = False
        method = "pdf2docx"

        # 1. Documents volumineux (> 50 pages) : PyMuPDF-Turbo immédiat
        if page_count > 50:
            method = "PyMuPDF-Turbo"
            try:
                convert_with_pymupdf_fast(pdf_p, docx_p)
                success = os.path.exists(to_long_path(docx_p)) and os.path.getsize(to_long_path(docx_p)) > 100
            except Exception:
                success = False
        else:
            # 2. Documents classiques (<= 50 pages) : pdf2docx avec timeout de 50s
            with ProcessPoolExecutor(max_workers=1) as executor:
                future = executor.submit(convert_with_pdf2docx_worker, str(pdf_p), str(docx_p))
                try:
                    ok, _ = future.result(timeout=50)
                    if ok and os.path.exists(to_long_path(docx_p)) and os.path.getsize(to_long_path(docx_p)) > 300:
                        success = True
                except Exception:
                    success = False
                finally:
                    executor.shutdown(wait=False, cancel_futures=True)

            # 3. Fallback immédiat si échec ou timeout de pdf2docx
            if not success:
                method = "PyMuPDF-Fallback"
                try:
                    convert_with_pymupdf_fast(pdf_p, docx_p)
                    success = os.path.exists(to_long_path(docx_p)) and os.path.getsize(to_long_path(docx_p)) > 100
                except Exception:
                    success = False

        dur = time.time() - t0
        if success:
            print(f"✅ OK [{method}] ({dur:.1f}s)")
            success_count += 1
            progress_data["success"] = success_count
        else:
            print(f"❌ ÉCHEC ({dur:.1f}s)")
            progress_data["errors"] = progress_data.get("errors", 0) + 1

        update_progress(progress_data)

    total_time = time.time() - start_time
    progress_data["status"] = "completed"
    progress_data["end_time"] = time.strftime("%Y-%m-%d %H:%M:%S")
    update_progress(progress_data)

    print("\n" + "=" * 75)
    print(f"🏁 CONVERSION TERMINÉE en {total_time/60:.1f} minutes")
    print(f"   ✅ Nouveaux DOCX créés avec succès : {success_count} / {total}")
    print("   📂 Les fichiers PDF et DOCX sont tous les deux disponibles dans vos dossiers.")
    print("=" * 75)

if __name__ == "__main__":
    main()
